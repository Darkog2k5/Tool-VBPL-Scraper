"""
Chuyển đổi HTML văn bản pháp luật sang DOCX giữ nguyên formatting gốc.

Giữ lại: font-weight (bold), font-style (italic), text-align (căn lề),
          font-size, table layout (header 2 cột), line break.
"""
import re
import unicodedata
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from loguru import logger

from scraper.detail_scraper import VBDocument


# ── CSS helpers ──────────────────────────────────────────────────────────────


def _parse_css(style_str: str | None) -> dict:
    if not style_str:
        return {}
    styles = {}
    for part in style_str.split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            styles[k.strip().lower()] = v.strip()
    return styles


def _get_alignment(css: dict):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    mapping = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }
    return mapping.get(css.get("text-align", "").lower())


def _get_font_size_pt(css: dict) -> float | None:
    m = re.match(r"([\d.]+)\s*pt", css.get("font-size", ""), re.IGNORECASE)
    return float(m.group(1)) if m else None


def _is_bold(css: dict) -> bool:
    return css.get("font-weight", "").lower() in ("bold", "700", "800", "900")


def _is_italic(css: dict) -> bool:
    return css.get("font-style", "").lower() == "italic"


def _nfc(text: str) -> str:
    return unicodedata.normalize("NFC", text) if text else ""


# ── HTML → DOCX runs ────────────────────────────────────────────────────────


def _add_runs(paragraph, element, parent_css: dict | None = None):
    """Đệ quy: chuyển children của element thành DOCX runs trong paragraph."""
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from lxml import etree

    if parent_css is None:
        parent_css = {}

    for child in element.children:
        if isinstance(child, NavigableString):
            raw = str(child).replace("\xa0", " ")
            text = _nfc(raw)
            if not text or (not text.strip() and text != " "):
                continue
            run = paragraph.add_run(text)
            if _is_bold(parent_css):
                run.bold = True
            if _is_italic(parent_css):
                run.italic = True
            sz = _get_font_size_pt(parent_css)
            if sz:
                run.font.size = Pt(sz)

        elif isinstance(child, Tag):
            if child.name == "br":
                # Line break trong cùng paragraph
                run = paragraph.add_run()
                br_elem = etree.SubElement(run._r, qn("w:br"))
                continue

            # Merge CSS cha + con
            child_css = {**parent_css, **_parse_css(child.get("style"))}
            if child.name in ("strong", "b"):
                child_css["font-weight"] = "bold"
            elif child.name in ("em", "i"):
                child_css["font-style"] = "italic"

            _add_runs(paragraph, child, child_css)


# ── HTML block → DOCX paragraph ─────────────────────────────────────────────


def _add_paragraph(doc, p_tag):
    """Chuyển <p> sang DOCX paragraph giữ alignment + runs."""
    css = _parse_css(p_tag.get("style"))
    text = p_tag.get_text(strip=True)

    if not text:
        doc.add_paragraph()          # Spacer
        return

    para = doc.add_paragraph()
    alignment = _get_alignment(css)
    if alignment is not None:
        para.alignment = alignment
    _add_runs(para, p_tag, css)


# ── HTML table → DOCX table ─────────────────────────────────────────────────


def _add_table(doc, table_tag):
    """Chuyển <table> sang DOCX table (giữ layout 2 cột header)."""
    from docx.oxml.ns import qn
    from lxml import etree

    rows = table_tag.find_all("tr")
    if not rows:
        return
    max_cols = max(len(r.find_all(["td", "th"])) for r in rows)
    if max_cols == 0:
        return

    tbl = doc.add_table(rows=len(rows), cols=max_cols)

    # Ẩn viền cho layout table (prov-table)
    is_layout = ("prov-table" in (table_tag.get("class") or [])
                 or "border: none" in (table_tag.get("style") or ""))
    if is_layout:
        tbl_pr = tbl._tbl.tblPr
        if tbl_pr is None:
            tbl_pr = etree.SubElement(tbl._tbl, qn("w:tblPr"))
        # Xoá borders cũ
        old = tbl_pr.find(qn("w:tblBorders"))
        if old is not None:
            tbl_pr.remove(old)
        borders = etree.SubElement(tbl_pr, qn("w:tblBorders"))
        for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = etree.SubElement(borders, qn(f"w:{name}"))
            b.set(qn("w:val"), "none")
            b.set(qn("w:sz"), "0")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), "auto")

    for i, row in enumerate(rows):
        cells = row.find_all(["td", "th"])
        for j, cell in enumerate(cells):
            if j >= max_cols:
                break
            docx_cell = tbl.cell(i, j)
            docx_cell.paragraphs[0].text = ""

            p_tags = cell.find_all("p")
            if p_tags:
                for k, p_tag in enumerate(p_tags):
                    para = docx_cell.paragraphs[0] if k == 0 else docx_cell.add_paragraph()
                    css = _parse_css(p_tag.get("style"))
                    alignment = _get_alignment(css)
                    if alignment is not None:
                        para.alignment = alignment
                    _add_runs(para, p_tag, css)
            else:
                text = _nfc(cell.get_text(strip=True))
                if text:
                    docx_cell.paragraphs[0].text = text


# ── Public API ───────────────────────────────────────────────────────────────


def convert_to_docx(vb: VBDocument, out_dir: Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    file_path = out_dir / "noi_dung.docx"
    doc = Document()

    # ── Default style ────────────────────────────────────────────────────────
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)

    content_html = getattr(vb, "content_html", "") or ""

    if content_html:
        # ── Parse HTML gốc, giữ nguyên formatting ───────────────────────────
        soup = BeautifulSoup(content_html, "html.parser")
        body = soup.find("body") or soup

        for element in body.children:
            if isinstance(element, NavigableString):
                text = _nfc(str(element).strip())
                if text:
                    doc.add_paragraph(text)
            elif isinstance(element, Tag):
                if element.name == "table":
                    _add_table(doc, element)
                elif element.name == "p":
                    _add_paragraph(doc, element)
                elif element.name in ("div", "section", "article"):
                    for child in element.children:
                        if isinstance(child, Tag):
                            if child.name == "table":
                                _add_table(doc, child)
                            elif child.name == "p":
                                _add_paragraph(doc, child)
                # Skip <head>, <html> wrappers

    elif vb.full_text:
        # ── Fallback: plain text (PDF, OCR, summary) ────────────────────────
        for line in vb.full_text.splitlines():
            line = _nfc(line.strip())
            if line:
                doc.add_paragraph(line)
    else:
        doc.add_paragraph("(Không có nội dung)")

    doc.save(str(file_path))
    logger.info(f"Đã lưu DOCX nội dung: {file_path}")
    return file_path